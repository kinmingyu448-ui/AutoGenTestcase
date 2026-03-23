import os
import json
import tempfile
from flask import Flask, request, Response, render_template, stream_with_context
import anthropic
from docx import Document

app = Flask(__name__)

SYSTEM_PROMPT_PATH = os.path.join(os.path.dirname(__file__), "system_prompt.txt")


def load_system_prompt():
    with open(SYSTEM_PROMPT_PATH, "r", encoding="utf-8") as f:
        return f.read()


def extract_docx_text(file_path):
    """从 .docx 文件提取纯文本（段落 + 表格）"""
    doc = Document(file_path)
    parts = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/generate", methods=["POST"])
def generate():
    api_key = request.form.get("api_key", "").strip()
    if not api_key:
        return Response("error: 请输入 API Key", status=400)

    # 获取需求文本
    requirement_text = request.form.get("requirement_text", "").strip()

    # 处理上传文件
    uploaded_file = request.files.get("file")
    if uploaded_file and uploaded_file.filename:
        filename = uploaded_file.filename.lower()
        if filename.endswith(".docx"):
            # Windows 兼容：先关闭临时文件再操作，最后删除
            tmp_dir = tempfile.mkdtemp()
            tmp_path = os.path.join(tmp_dir, "upload.docx")
            try:
                uploaded_file.save(tmp_path)
                file_text = extract_docx_text(tmp_path)
            finally:
                try:
                    os.remove(tmp_path)
                    os.rmdir(tmp_dir)
                except OSError:
                    pass
            requirement_text = file_text if not requirement_text else requirement_text + "\n\n" + file_text
        elif filename.endswith(".txt"):
            file_text = uploaded_file.read().decode("utf-8")
            requirement_text = file_text if not requirement_text else requirement_text + "\n\n" + file_text

    if not requirement_text:
        return Response("error: 请输入需求文档内容或上传文件", status=400)

    # 获取技术文档
    tech_doc_text = request.form.get("tech_doc_text", "").strip()

    # 处理技术文档上传文件
    tech_file = request.files.get("tech_file")
    if tech_file and tech_file.filename:
        tech_filename = tech_file.filename.lower()
        if tech_filename.endswith(".docx"):
            tmp_dir = tempfile.mkdtemp()
            tmp_path = os.path.join(tmp_dir, "tech_upload.docx")
            try:
                tech_file.save(tmp_path)
                tech_file_text = extract_docx_text(tmp_path)
            finally:
                try:
                    os.remove(tmp_path)
                    os.rmdir(tmp_dir)
                except OSError:
                    pass
            tech_doc_text = tech_file_text if not tech_doc_text else tech_doc_text + "\n\n" + tech_file_text
        elif tech_filename.endswith(".txt"):
            tech_file_text = tech_file.read().decode("utf-8")
            tech_doc_text = tech_file_text if not tech_doc_text else tech_doc_text + "\n\n" + tech_file_text

    # 获取设计稿描述
    design_doc_text = request.form.get("design_doc_text", "").strip()

    # 处理设计稿图片（转为 base64 供视觉模型使用，支持多张）
    import base64
    design_images = []
    design_files = request.files.getlist("design_files")
    for df in design_files:
        if df and df.filename:
            img_bytes = df.read()
            img_b64 = base64.b64encode(img_bytes).decode("utf-8")
            ext = df.filename.rsplit(".", 1)[-1].lower()
            media_type = f"image/{'jpeg' if ext in ('jpg', 'jpeg') else 'png'}"
            design_images.append({"data": img_b64, "media_type": media_type})

    system_prompt = load_system_prompt()

    # 组装用户消息
    user_message_parts = [f"请根据以下需求文档生成测试用例：\n\n{requirement_text}"]
    if tech_doc_text:
        user_message_parts.append(f"\n\n---\n以下是技术文档，请结合技术实现细节生成更精准的测试用例：\n\n{tech_doc_text}")
    if design_doc_text:
        user_message_parts.append(f"\n\n---\n以下是设计稿描述，请结合UI/交互设计生成相关的测试用例：\n\n{design_doc_text}")

    user_message = "".join(user_message_parts)

    # 构建消息内容（支持多张图片）
    if design_images:
        user_content = [
            {"type": "text", "text": user_message},
            {"type": "text", "text": f"\n\n---\n以下是 {len(design_images)} 张设计稿图片，请结合设计稿的UI布局和交互细节生成相关测试用例："}
        ]
        for img in design_images:
            user_content.append({"type": "image", "source": {"type": "base64", "media_type": img["media_type"], "data": img["data"]}})
    else:
        user_content = user_message

    def stream_response():
        try:
            client = anthropic.Anthropic(api_key=api_key)
            with client.messages.stream(
                model="claude-sonnet-4-20250514",
                max_tokens=8192,
                system=system_prompt,
                messages=[{"role": "user", "content": user_content}],
            ) as stream:
                for text in stream.text_stream:
                    chunk = f"data: {json.dumps({'text': text}, ensure_ascii=False)}\n\n"
                    yield chunk.encode("utf-8")
            yield f"data: {json.dumps({'done': True})}\n\n".encode("utf-8")
        except anthropic.AuthenticationError:
            yield f"data: {json.dumps({'error': 'API Key 无效，请检查后重试'}, ensure_ascii=False)}\n\n".encode("utf-8")
        except anthropic.RateLimitError:
            yield f"data: {json.dumps({'error': 'API 请求频率超限，请稍后重试'}, ensure_ascii=False)}\n\n".encode("utf-8")
        except Exception as e:
            yield f"data: {json.dumps({'error': f'生成失败: {str(e)}'}, ensure_ascii=False)}\n\n".encode("utf-8")

    return Response(
        stream_with_context(stream_response()),
        content_type="text/event-stream; charset=utf-8",
        headers={
            "Cache-Control": "no-cache",
            "X-Accel-Buffering": "no",
        },
    )


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)
